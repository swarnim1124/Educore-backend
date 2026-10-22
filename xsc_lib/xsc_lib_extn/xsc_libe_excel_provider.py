"""
File: excel_provider.py
Description: External integration for Excel generation using openpyxl.
Author: Siddharth Amdal.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from typing import Dict, Any


class ExcelProvider:
    @staticmethod
    def generate(data: Dict[str, Any], output_path: str) -> bool:
        """
        Takes raw data payload, builds an Excel workbook, and saves it to output_path.

        Args:
            data (Dict[str, Any]): The dynamic data payload to inject.
            output_path (str): The destination file path.

        Returns:
            bool: True if generated successfully.
        """
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Attendance Report"

            # Styles
            header_font = Font(bold=True, size=14)
            label_font = Font(bold=True)
            center_align = Alignment(horizontal="center")
            fill_header = PatternFill(
                start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
            )

            # Title
            ws.merge_cells("A1:D1")
            title_cell = ws["A1"]
            title_cell.value = (
                f"Attendance Report - {data.get('institution_name', 'Institution')}"
            )
            title_cell.font = header_font
            title_cell.alignment = center_align

            # Student Info
            ws["A3"] = "Student Name:"
            ws["A3"].font = label_font
            ws["B3"] = data.get("student_name", "")

            ws["A4"] = "USN:"
            ws["A4"].font = label_font
            ws["B4"] = data.get("usn", "")

            ws["A5"] = "Semester:"
            ws["A5"].font = label_font
            ws["B5"] = data.get("semester", "")

            ws["A6"] = "Overall Percentage:"
            ws["A6"].font = label_font
            ws["B6"] = f"{data.get('overall_percentage', '')}%"

            # Table Headers
            headers = ["Subject", "Total Classes", "Attended", "Percentage"]
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=8, column=col)
                cell.value = header
                cell.font = label_font
                cell.fill = fill_header

            # Table Data
            row = 9
            for sub in data.get("subjects", []):
                ws.cell(row=row, column=1, value=sub.get("subject", ""))
                ws.cell(row=row, column=2, value=sub.get("total_classes", 0))
                ws.cell(row=row, column=3, value=sub.get("attended", 0))
                ws.cell(row=row, column=4, value=f"{sub.get('percentage', 0)}%")
                row += 1

            # Column widths
            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 15
            ws.column_dimensions["C"].width = 15
            ws.column_dimensions["D"].width = 15

            wb.save(output_path)
            return True
        except Exception as _e:
            print(f"[ExcelProvider] Error generating Excel: {str(e)}")
            raise
