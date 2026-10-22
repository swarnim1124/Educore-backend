"""
File: word_provider.py
Description: External integration for Word document generation using python-docx.
Author: Siddharth Amdal.
"""

from docx import Document
from typing import Dict, Any


class WordProvider:
    @staticmethod
    def generate(template_path: str, data: Dict[str, Any], output_path: str) -> bool:
        """
        Loads a docx template, replaces placeholders with data, and saves to output_path.

        Args:
            template_path (str): The physical path to the docx template file.
            data (Dict[str, Any]): The dynamic data payload to inject.
            output_path (str): The destination file path.

        Returns:
            bool: True if generated successfully.
        """
        try:
            doc = Document(template_path)

            # Simple paragraph placeholder replacement
            for p in doc.paragraphs:
                for key, value in data.items():
                    # Handle both spaced and unspaced placeholders just in case
                    placeholder_spaced = f"{{{{ {key} }}}}"
                    placeholder_tight = f"{{{{{key}}}}}"
                    if placeholder_spaced in p.text:
                        p.text = p.text.replace(placeholder_spaced, str(value))
                    if placeholder_tight in p.text:
                        p.text = p.text.replace(placeholder_tight, str(value))

            # Simple table placeholder replacement (often used for nested data like subjects)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for key, value in data.items():
                                placeholder_spaced = f"{{{{ {key} }}}}"
                                placeholder_tight = f"{{{{{key}}}}}"
                                if placeholder_spaced in p.text:
                                    p.text = p.text.replace(
                                        placeholder_spaced, str(value)
                                    )
                                if placeholder_tight in p.text:
                                    p.text = p.text.replace(
                                        placeholder_tight, str(value)
                                    )

            doc.save(output_path)
            return True
        except Exception as _e:
            print(f"[WordProvider] Error generating Word document: {str(e)}")
            raise
